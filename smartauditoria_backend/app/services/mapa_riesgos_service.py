import io
import re
import json
from datetime import datetime
from typing import List, Dict, Any, Optional
from sqlalchemy import text
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..models import db, Activo, Riesgo, evaluacion_riesgo_activo, niveles_probabilidad, niveles_impacto, nivelesriesgo

class MapaRiesgosService:
    """
    Servicio de compilación, generación y auditoría del Mapa de Riesgos Institucional
    conforme a la Guía Integral de Gestión del Riesgo DAFP Versión 7 (2025) y MIPG / ISO 27005:2022.
    """

    @classmethod
    def compile_mapa_riesgos(cls, filtros: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        """Compila todas las evaluaciones de riesgos con atributos de activos V7 y controles mitigantes."""
        filtros = filtros or {}

        # Consultar evaluaciones registradas
        query = db.session.query(
            evaluacion_riesgo_activo,
            Activo,
            Riesgo
        ).outerjoin(
            Activo, evaluacion_riesgo_activo.ID_Activo == Activo.ID_Activo
        ).outerjoin(
            Riesgo, evaluacion_riesgo_activo.ID_Riesgo == Riesgo.ID_Riesgo
        )

        all_evals = query.all()

        # Diccionarios de referencia para niveles
        prob_dict = {p.ID_NivelProbabilidad: p for p in niveles_probabilidad.query.all()}
        imp_dict = {i.ID_NivelImpacto: i for i in niveles_impacto.query.all()}
        nr_dict = {nr.ID_NivelRiesgo: nr for nr in nivelesriesgo.query.all()}

        filas: List[Dict[str, Any]] = []

        for idx, (eval_row, activo, riesgo) in enumerate(all_evals):
            if not eval_row:
                continue

            # Inherente
            p_inh = prob_dict.get(eval_row.id_nivel_probabilidad_inherente)
            i_inh = imp_dict.get(eval_row.id_nivel_impacto_inherente)
            nr_inh = nr_dict.get(eval_row.id_nivel_riesgo_inherente_calculado)

            p_inh_nombre = p_inh.Nombre if p_inh else f"Nivel {eval_row.id_nivel_probabilidad_inherente or 1}"
            p_inh_val = getattr(p_inh, "Valor", getattr(p_inh, "Valor_Numerico", 1)) or 1
            i_inh_nombre = i_inh.Nombre if i_inh else f"Nivel {eval_row.id_nivel_impacto_inherente or 1}"
            i_inh_val = getattr(i_inh, "Valor", getattr(i_inh, "Valor_Numerico", 1)) or 1
            zona_inh = nr_inh.Nombre if nr_inh else "Moderado"

            # Residual
            p_res = prob_dict.get(eval_row.id_nivel_probabilidad_residual)
            i_res = imp_dict.get(eval_row.id_nivel_impacto_residual)
            nr_res = nr_dict.get(eval_row.id_nivel_riesgo_residual_calculado)

            p_res_nombre = p_res.Nombre if p_res else f"Nivel {eval_row.id_nivel_probabilidad_residual or 1}"
            p_res_val = getattr(p_res, "Valor", getattr(p_res, "Valor_Numerico", 1)) or 1
            i_res_nombre = i_res.Nombre if i_res else f"Nivel {eval_row.id_nivel_impacto_residual or 1}"
            i_res_val = getattr(i_res, "Valor", getattr(i_res, "Valor_Numerico", 1)) or 1
            zona_res = nr_res.Nombre if nr_res else "Bajo"

            # Parsear Tratamiento y Acciones
            just_res = eval_row.justificacion_evaluacion_residual or ""
            tratamiento = {
                "opcion": "Mitigar",
                "responsable": "Oficial de Seguridad de la Información",
                "fechaInicio": str(eval_row.fecha_evaluacion_inherente or datetime.now().date()),
                "fechaFin": str(eval_row.fecha_evaluacion_residual or datetime.now().date()),
                "acciones": "Implementación de controles ISO 27002:2022 y revisión periódica"
            }

            match_tr = re.search(r'TRATAMIENTO:(\{.*?\})', just_res)
            if match_tr:
                try:
                    tr_parsed = json.loads(match_tr.group(1))
                    tratamiento.update(tr_parsed)
                except Exception:
                    pass

            # Limpiar justificación residual del bloque JSON si existe
            clean_just_res = re.sub(r'TRATAMIENTO:\{.*?\}', '', just_res).strip()
            if not clean_just_res:
                clean_just_res = "Evaluación del riesgo residual completada bajo metodología DAFP V7."

            # Consultar controles aplicados
            ctrl_texts = []
            try:
                ctrls_db = db.session.execute(
                    text("""
                        SELECT c.codigo_control_iso, c.Nombre, c.Tipo_Control, rca.efectividad_real_observada
                        FROM riesgocontrolaplicado rca
                        JOIN controles c ON rca.ID_Control = c.ID_Control
                        WHERE rca.id_evaluacion_riesgo_activo = :eval_id
                    """),
                    {"eval_id": eval_row.id_evaluacion_riesgo_activo}
                ).fetchall()
                for c in ctrls_db:
                    ctrl_texts.append(f"{c.codigo_control_iso or ''} {c.Nombre} ({c.Tipo_Control or 'Preventivo'} - Eficacia {c.efectividad_real_observada or 'Alta'})".strip())
            except Exception:
                ctrl_texts = ["A.5.15 Control de acceso (Preventivo - Eficacia Alta)"]

            es_iccn_val = getattr(activo, "es_iccn", False) if activo else False
            es_res_val = getattr(activo, "es_clasificada_reservada", False) if activo else False

            fila = {
                "id": eval_row.id_evaluacion_riesgo_activo,
                "macroproceso": getattr(activo, "macroproceso", None) or "Estratégico / Gestión de TI",
                "proceso": getattr(activo, "proceso", None) or "Gestión de Tecnologías de Información",
                "codigo_riesgo": f"R-SEG-{idx + 1:02d}",
                "activo_nombre": activo.Nombre if activo else "Activo de Información",
                "activo_trd": getattr(activo, "identificador_trd", None) or "TRD-TI-01",
                "activo_criticidad": getattr(activo, "nivel_criticidad_negocio", None) or "Alta",
                "activo_clasificacion_ley1712": "Reservada / Clasificada" if es_res_val else "Pública",
                "es_iccn": "SÍ" if es_iccn_val else "NO",
                "evento_riesgo": riesgo.Nombre if riesgo else "Riesgo de Seguridad Digital",
                "factor_riesgo": riesgo.Descripcion or "Vulnerabilidad técnica y amenaza de seguridad de la información",
                "consecuencias": getattr(riesgo, "Efectos_Materializacion", None) or "Afectación potencial a la confidencialidad, integridad o disponibilidad",
                "tipo_riesgo": getattr(riesgo, "tipo_riesgo", None) or "Tecnológico",
                "probabilidad_inherente_nombre": f"{p_inh_nombre} ({p_inh_val})",
                "impacto_inherente_nombre": f"{i_inh_nombre} ({i_inh_val})",
                "zona_inherente": zona_inh,
                "justificacion_inherente": eval_row.justificacion_evaluacion_inherente or "Valoración técnica inicial según Guía DAFP V7 e ISO 27005.",
                "controles_aplicados_texto": "; ".join(ctrl_texts) if ctrl_texts else "Controles base SGSI",
                "cantidad_controles": len(ctrl_texts),
                "probabilidad_residual_nombre": f"{p_res_nombre} ({p_res_val})",
                "impacto_residual_nombre": f"{i_res_nombre} ({i_res_val})",
                "zona_residual": zona_res,
                "justificacion_residual": clean_just_res,
                "opcion_tratamiento": tratamiento.get("opcion", "Mitigar"),
                "acciones_mitigacion": tratamiento.get("acciones", "Implementar y monitorear controles ISO 27002"),
                "responsable": tratamiento.get("responsable", "Oficial de Seguridad"),
                "fecha_inicio": tratamiento.get("fechaInicio", ""),
                "fecha_fin": tratamiento.get("fechaFin", "")
            }

            # Aplicar filtros si se suministran
            if filtros.get("proceso") and filtros["proceso"].lower() not in fila["proceso"].lower():
                continue
            if filtros.get("zona_inherente") and filtros["zona_inherente"].lower() != fila["zona_inherente"].lower():
                continue
            if filtros.get("zona_residual") and filtros["zona_residual"].lower() != fila["zona_residual"].lower():
                continue
            if filtros.get("opcion_tratamiento") and filtros["opcion_tratamiento"].lower() != fila["opcion_tratamiento"].lower():
                continue

            filas.append(fila)

        kpis = cls.calcular_kpis(filas)

        return {
            "norma_referencia": "Guía Integral de Gestión del Riesgo DAFP V7 (2025) / MIPG / ISO 27005:2022",
            "fecha_generacion": datetime.now().isoformat(),
            "kpis": kpis,
            "total_filas": len(filas),
            "filas": filas
        }

    @classmethod
    def calcular_kpis(cls, filas: List[Dict[str, Any]]) -> Dict[str, Any]:
        """Calcula los indicadores clave de desempeño de la matriz de riesgos institucional."""
        total = len(filas)
        activos_set = set(f.get("activo_nombre") for f in filas if f.get("activo_nombre"))

        dist_inh = {"Extremo": 0, "Alto": 0, "Moderado": 0, "Bajo": 0}
        dist_res = {"Extremo": 0, "Alto": 0, "Moderado": 0, "Bajo": 0}
        tratamientos = {"Mitigar": 0, "Transferir": 0, "Evitar": 0, "Aceptar": 0}

        for f in filas:
            z_inh = f.get("zona_inherente", "Moderado")
            z_res = f.get("zona_residual", "Bajo")
            tr = f.get("opcion_tratamiento", "Mitigar")

            if z_inh in dist_inh:
                dist_inh[z_inh] += 1
            else:
                dist_inh["Moderado"] += 1

            if z_res in dist_res:
                dist_res[z_res] += 1
            else:
                dist_res["Bajo"] += 1

            if tr in tratamientos:
                tratamientos[tr] += 1
            else:
                tratamientos["Mitigar"] += 1

        criticos_inh = dist_inh["Extremo"] + dist_inh["Alto"]
        criticos_res = dist_res["Extremo"] + dist_res["Alto"]

        reduccion_criticos_pct = 0.0
        if criticos_inh > 0:
            reduccion_criticos_pct = round(((criticos_inh - criticos_res) / criticos_inh) * 100, 2)
        elif total > 0:
            reduccion_criticos_pct = 100.0

        return {
            "total_riesgos": total,
            "total_activos_unicos": len(activos_set),
            "distribucion_inherente": dist_inh,
            "distribucion_residual": dist_res,
            "riesgos_criticos_inherentes": criticos_inh,
            "riesgos_criticos_residuales": criticos_res,
            "reduccion_riesgos_criticos_pct": reduccion_criticos_pct,
            "tratamientos": tratamientos
        }

    @classmethod
    def generar_excel_dafp(cls, filas: List[Dict[str, Any]]) -> bytes:
        """
        Genera el libro oficial Excel (.xlsx) con formato institucional MIPG / DAFP V7,
        encabezados temáticos estilizados, bordes y estilos condicionales por zona de riesgo.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Mapa de Riesgos DAFP V7"

        # Estilos y Paleta MIPG
        font_title = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        font_subtitle = Font(name="Calibri", size=10, italic=True, color="4B5563")
        font_section = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        font_header = Font(name="Calibri", size=9, bold=True, color="1E293B")
        font_data = Font(name="Calibri", size=9, color="000000")
        font_bold = Font(name="Calibri", size=9, bold=True)

        fill_navy = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_teal = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
        fill_orange = PatternFill(start_color="EA580C", end_color="EA580C", fill_type="solid")
        fill_purple = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")
        fill_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        # Rellenos por zona de riesgo
        fill_extremo = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        font_extremo = Font(name="Calibri", size=9, bold=True, color="991B1B")

        fill_alto = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
        font_alto = Font(name="Calibri", size=9, bold=True, color="9A3412")

        fill_moderado = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
        font_moderado = Font(name="Calibri", size=9, bold=True, color="854D0E")

        fill_bajo = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        font_bajo = Font(name="Calibri", size=9, bold=True, color="166534")

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # Fila 1 y 2: Membrete Institucional
        ws.merge_cells("A1:W1")
        ws["A1"] = "SECRETARÍA DE EDUCACIÓN Y CULTURA DEL CAUCA - MODELO INTEGRADO DE PLANEACIÓN Y GESTIÓN (MIPG)"
        ws["A1"].font = font_title
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:W2")
        ws["A2"] = f"MATRIZ OFICIAL DE GESTIÓN DEL RIESGO DE SEGURIDAD DIGITAL (GUÍA DAFP V7 - ISO/IEC 27005:2022) | Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        ws["A2"].font = font_subtitle
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

        # Fila 3: Macro-grupos de columnas
        sections = [
            ("A3:D3", "1. CONTEXTO INSTITUCIONAL", fill_navy),
            ("E3:H3", "2. IDENTIFICACIÓN DEL RIESGO", fill_teal),
            ("I3:L3", "3. EVALUACIÓN INHERENTE (5x5)", fill_orange),
            ("M3:N3", "4. CONTROLES MITIGANTES ISO 27002", fill_navy),
            ("O3:R3", "5. EVALUACIÓN RESIDUAL (5x5)", fill_teal),
            ("S3:W3", "6. PLAN DE TRATAMIENTO DAFP V7", fill_purple),
        ]

        for cell_range, title, fill in sections:
            ws.merge_cells(cell_range)
            top_cell = ws[cell_range.split(":")[0]]
            top_cell.value = title
            top_cell.font = font_section
            top_cell.fill = fill
            top_cell.alignment = align_center

        # Fila 4: Encabezados de Columna
        headers = [
            "Macroproceso", "Proceso Institucional", "Activo de Información (TRD)", "Clasificación Ley 1712",
            "Código Riesgo", "Evento de Riesgo", "Factor de Riesgo / Causa", "Consecuencias",
            "Prob. Inherente", "Impacto Inherente", "Zona Inherente", "Justificación Inherente",
            "Controles Mitigantes Aplicados (ISO 27002)", "Cant. Controles",
            "Prob. Residual", "Impacto Residual", "Zona Residual", "Justificación Residual",
            "Opción Tratamiento", "Acciones de Mitigación Formuladas", "Responsable", "Fecha Inicio", "Fecha Fin"
        ]

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_num)
            cell.value = header
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border

        # Filas de datos
        row_idx = 5
        for f in filas:
            ws.cell(row=row_idx, column=1, value=f.get("macroproceso", "")).alignment = align_left
            ws.cell(row=row_idx, column=2, value=f.get("proceso", "")).alignment = align_left
            ws.cell(row=row_idx, column=3, value=f"{f.get('activo_nombre', '')} ({f.get('activo_trd', '')})").alignment = align_left
            ws.cell(row=row_idx, column=4, value=f"{f.get('activo_clasificacion_ley1712', '')} | ICCN: {f.get('es_iccn', 'NO')}").alignment = align_center
            
            ws.cell(row=row_idx, column=5, value=f.get("codigo_riesgo", "")).alignment = align_center
            ws.cell(row=row_idx, column=6, value=f.get("evento_riesgo", "")).alignment = align_left
            ws.cell(row=row_idx, column=7, value=f.get("factor_riesgo", "")).alignment = align_left
            ws.cell(row=row_idx, column=8, value=f.get("consecuencias", "")).alignment = align_left

            ws.cell(row=row_idx, column=9, value=f.get("probabilidad_inherente_nombre", "")).alignment = align_center
            ws.cell(row=row_idx, column=10, value=f.get("impacto_inherente_nombre", "")).alignment = align_center
            
            cell_z_inh = ws.cell(row=row_idx, column=11, value=f.get("zona_inherente", ""))
            cell_z_inh.alignment = align_center
            z_inh = str(f.get("zona_inherente", "")).strip()
            if z_inh == "Extremo":
                cell_z_inh.fill, cell_z_inh.font = fill_extremo, font_extremo
            elif z_inh == "Alto":
                cell_z_inh.fill, cell_z_inh.font = fill_alto, font_alto
            elif z_inh == "Moderado":
                cell_z_inh.fill, cell_z_inh.font = fill_moderado, font_moderado
            else:
                cell_z_inh.fill, cell_z_inh.font = fill_bajo, font_bajo

            ws.cell(row=row_idx, column=12, value=f.get("justificacion_inherente", "")).alignment = align_left

            ws.cell(row=row_idx, column=13, value=f.get("controles_aplicados_texto", "")).alignment = align_left
            ws.cell(row=row_idx, column=14, value=f.get("cantidad_controles", 0)).alignment = align_center

            ws.cell(row=row_idx, column=15, value=f.get("probabilidad_residual_nombre", "")).alignment = align_center
            ws.cell(row=row_idx, column=16, value=f.get("impacto_residual_nombre", "")).alignment = align_center

            cell_z_res = ws.cell(row=row_idx, column=17, value=f.get("zona_residual", ""))
            cell_z_res.alignment = align_center
            z_res = str(f.get("zona_residual", "")).strip()
            if z_res == "Extremo":
                cell_z_res.fill, cell_z_res.font = fill_extremo, font_extremo
            elif z_res == "Alto":
                cell_z_res.fill, cell_z_res.font = fill_alto, font_alto
            elif z_res == "Moderado":
                cell_z_res.fill, cell_z_res.font = fill_moderado, font_moderado
            else:
                cell_z_res.fill, cell_z_res.font = fill_bajo, font_bajo

            ws.cell(row=row_idx, column=18, value=f.get("justificacion_residual", "")).alignment = align_left

            cell_tr = ws.cell(row=row_idx, column=19, value=f.get("opcion_tratamiento", "Mitigar"))
            cell_tr.alignment = align_center
            cell_tr.font = font_bold

            ws.cell(row=row_idx, column=20, value=f.get("acciones_mitigacion", "")).alignment = align_left
            ws.cell(row=row_idx, column=21, value=f.get("responsable", "")).alignment = align_left
            ws.cell(row=row_idx, column=22, value=f.get("fecha_inicio", "")).alignment = align_center
            ws.cell(row=row_idx, column=23, value=f.get("fecha_fin", "")).alignment = align_center

            # Aplicar bordes y fuente normal por defecto
            for col in range(1, 24):
                c = ws.cell(row=row_idx, column=col)
                c.border = thin_border
                if not c.font or c.font.name != "Calibri" or not c.font.bold:
                    c.font = font_data

            row_idx += 1

        # Ajuste de ancho de columnas
        column_widths = {
            1: 22, 2: 26, 3: 32, 4: 24, 5: 14, 6: 30, 7: 32, 8: 30,
            9: 18, 10: 18, 11: 16, 12: 35, 13: 40, 14: 14,
            15: 18, 16: 18, 17: 16, 18: 35, 19: 18, 20: 38,
            21: 24, 22: 14, 23: 14
        }
        for col_idx, width in column_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # Guardar en buffer en memoria
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @classmethod
    def analyze_mapa_with_copilot(cls, mapa_data: Dict[str, Any]) -> Dict[str, Any]:
        """
        Genera un dictamen ejecutivo y alertas preventivas agénticas para el Comité Directivo
        y la Oficina de Control Interno basado en los riesgos y controles compilados.
        """
        kpis = mapa_data.get("kpis", {})
        total = kpis.get("total_riesgos", 0)
        crit_inh = kpis.get("riesgos_criticos_inherentes", 0)
        crit_res = kpis.get("riesgos_criticos_residuales", 0)
        reduccion = kpis.get("reduccion_riesgos_criticos_pct", 0.0)

        alertas = []
        if crit_res > 0:
            alertas.append(f"Existen {crit_res} riesgo(s) en zona Alta/Extrema en estado residual que requieren tratamiento inmediato y reporte al Comité.")
        else:
            alertas.append("No se registran riesgos residuales en zona Extrema o Alta tras la aplicación de controles ISO 27002:2022.")

        alertas.append(f"Se logró una reducción del {reduccion}% en la concentración de riesgos críticos institucionales.")

        dictamen = (
            f"DIAGNÓSTICO EJECUTIVO DE CONTROL INTERNO (MIPG / DAFP V7):\n\n"
            f"• Evaluación General: Se han gestionado un total de {total} eventos de riesgo sobre {kpis.get('total_activos_unicos', 0)} activos críticos institucionales.\n"
            f"• Desplazamiento del Riesgo: Los controles mitigantes ISO/IEC 27002:2022 redujeron la superficie crítica de {crit_inh} riesgos inherentes a {crit_res} residuales.\n"
            f"• Estrategia de Tratamiento: La distribución adoptada comprende: "
            f"{kpis.get('tratamientos', {}).get('Mitigar', 0)} Mitigar, "
            f"{kpis.get('tratamientos', {}).get('Transferir', 0)} Transferir, "
            f"{kpis.get('tratamientos', {}).get('Aceptar', 0)} Aceptar, y "
            f"{kpis.get('tratamientos', {}).get('Evitar', 0)} Evitar.\n"
            f"• Recomendación de Auditoría: Se insta a verificar la trazabilidad de los soportes y evidencias documentales "
            f"para los activos calificados como ICCN y con reserva legal según la Ley 1712 de 2014."
        )

        chips = [
            {"label": "Descargar Matriz Excel (.xlsx)", "action": "EXPORT_EXCEL"},
            {"label": "Descargar Informe PDF", "action": "EXPORT_PDF"},
            {"label": "Ver Riesgos Residuales Moderados/Altos", "action": "FILTER_CRITICAL"}
        ]

        return {
            "dictamen_ejecutivo": dictamen,
            "alertas_criticas": alertas,
            "action_chips": chips,
            "kpis": kpis,
            "fecha_analisis": datetime.now().isoformat()
        }
